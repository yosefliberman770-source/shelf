#!/usr/bin/env python3
"""Inspect a candidate's real files: list them (Zenodo, Harvard Dataverse, figshare), fetch small ones, report fields.

  python3 scripts/historical-data/discover/inspect_sources.py <doi> [<doi> ...]
Results are appended to data/historical/discovery/inspected.jsonl (one line per candidate) and the files kept in
scripts/atlas-build/.cache/inspect/<doi>/.
"""
import csv
import io
import json
import os
import re
import sys
import urllib.request
import zipfile

ROOT = os.path.join(os.path.dirname(__file__), '..', '..', '..')
CACHE = os.path.join(ROOT, 'scripts', 'atlas-build', '.cache', 'inspect')
OUT = os.path.join(ROOT, 'data', 'historical', 'discovery', 'inspected.jsonl')
UA = {'User-Agent': 'ShelfResearch/1.0 (personal historical atlas)'}
MAX = int(os.environ.get('INSPECT_MAX_MB', '80')) * 1024 * 1024


def get(url, binary=False):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=180) as r:
        data = r.read()
    return data if binary else json.loads(data)


def files_for(doi):
    d = doi.lower()
    if 'zenodo' in d:
        rec = get(f'https://zenodo.org/api/records/{d.rsplit(".", 1)[-1]}')
        meta = rec['metadata']
        return meta, [{'name': f['key'], 'size': f['size'], 'url': f['links']['self']} for f in rec.get('files', [])]
    if d.startswith('10.7910/'):
        rec = get(f'https://dataverse.harvard.edu/api/datasets/:persistentId/?persistentId=doi:{doi}')
        v = rec['data']['latestVersion']
        meta = {f['typeName']: f['value'] for f in v['metadataBlocks']['citation']['fields']}
        return meta, [{'name': f['dataFile']['filename'], 'size': f['dataFile'].get('filesize'), 'url': f"https://dataverse.harvard.edu/api/access/datafile/{f['dataFile']['id']}?format=original",
                       'restricted': f.get('restricted')} for f in v['files']]
    if 'figshare' in d:
        aid = re.search(r'figshare\.(\d+)', d).group(1)
        rec = get(f'https://api.figshare.com/v2/articles/{aid}')
        return {'title': rec['title'], 'license': rec.get('license', {}).get('name')}, [{'name': f['name'], 'size': f['size'], 'url': f['download_url']} for f in rec['files']]
    raise ValueError('unsupported repository: ' + doi)


def describe_table(name, data):
    text = data.decode('utf-8-sig', errors='replace')
    sep = max([',', ';', '\t', '|'], key=lambda s: text[:5000].count(s))
    rows = list(csv.reader(io.StringIO(text), delimiter=sep))
    head = rows[0] if rows else []
    return {'file': name, 'rows': len(rows) - 1, 'columns': head[:60], 'sample': rows[1:3]}


def describe(name, data):
    low = name.lower()
    if low.endswith(('.csv', '.tsv', '.txt', '.tab')):
        return [describe_table(name, data)]
    if low.endswith(('.geojson', '.json')):
        try:
            d = json.loads(data)
        except Exception:  # noqa: BLE001
            return [{'file': name, 'note': 'not JSON'}]
        feats = d.get('features') if isinstance(d, dict) else None
        if feats:
            return [{'file': name, 'features': len(feats), 'geometry': sorted({(f.get('geometry') or {}).get('type', '') for f in feats[:500]}),
                     'properties': list((feats[0].get('properties') or {}).keys())[:60], 'sample': [f.get('properties') for f in feats[:2]]}]
        return [{'file': name, 'jsonKeys': list(d)[:30] if isinstance(d, dict) else f'list of {len(d)}', 'sample': (d[:2] if isinstance(d, list) else None)}]
    if low.endswith('.zip'):
        out = []
        z = zipfile.ZipFile(io.BytesIO(data))
        for n in z.namelist()[:200]:
            if n.lower().endswith(('.csv', '.tsv', '.txt', '.geojson', '.json')) and z.getinfo(n).file_size < MAX:
                out += describe(n, z.read(n))
            elif n.lower().endswith('.dbf'):
                try:
                    import shapefile
                    r = shapefile.Reader(dbf=io.BytesIO(z.read(n)), encoding='utf-8', encodingErrors='replace')
                    out.append({'file': n, 'rows': len(r), 'columns': [f[0] for f in r.fields[1:]][:60], 'sample': [list(r.record(i)) for i in range(min(2, len(r)))]})
                except Exception as e:  # noqa: BLE001
                    out.append({'file': n, 'note': str(e)[:100]})
        out.append({'zipMembers': len(z.namelist()), 'firstMembers': z.namelist()[:25]})
        return out
    if low.endswith('.xlsx'):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
        out = []
        for ws in wb.worksheets[:5]:
            rows = [list(r) for _, r in zip(range(4), ws.iter_rows(values_only=True))]
            out.append({'file': name, 'sheet': ws.title, 'rows': ws.max_row, 'columns': rows[0] if rows else [], 'sample': rows[1:3]})
        return out
    if low.endswith('.dbf'):
        import shapefile
        r = shapefile.Reader(dbf=io.BytesIO(data), encoding='utf-8', encodingErrors='replace')
        return [{'file': name, 'rows': len(r), 'columns': [f[0] for f in r.fields[1:]], 'sample': [list(r.record(i)) for i in range(min(2, len(r)))]}]
    return [{'file': name, 'note': 'not inspected (format)'}]


def inspect(doi):
    meta, files = files_for(doi)
    d = os.path.join(CACHE, re.sub(r'[^\w.-]', '_', doi))
    os.makedirs(d, exist_ok=True)
    rep = {'doi': doi, 'title': meta.get('title'), 'licence': (meta.get('license') or {}).get('id') if isinstance(meta.get('license'), dict) else meta.get('license'),
           'files': [{k: f[k] for k in ('name', 'size') if k in f} | ({'restricted': True} if f.get('restricted') else {}) for f in files], 'contents': []}
    for f in files:
        if f.get('restricted') or (f.get('size') or 0) > MAX:
            continue
        if not re.search(r'\.(csv|tsv|txt|tab|json|geojson|zip|dbf|xlsx)$', f['name'].lower()):
            continue
        p = os.path.join(d, f['name'].replace('/', '__'))
        if not os.path.exists(p):
            open(p, 'wb').write(get(f['url'], binary=True))
        try:
            rep['contents'] += describe(f['name'], open(p, 'rb').read())
        except Exception as e:  # noqa: BLE001
            rep['contents'].append({'file': f['name'], 'error': str(e)[:200]})
    with open(OUT, 'a', encoding='utf-8') as fh:
        fh.write(json.dumps(rep, ensure_ascii=False, default=str) + '\n')
    return rep


if __name__ == '__main__':
    for doi in sys.argv[1:]:
        try:
            r = inspect(doi)
            print(json.dumps(r, ensure_ascii=False, default=str)[:2500])
        except Exception as e:  # noqa: BLE001
            print(doi, 'FAILED', e)
        print('=' * 80)
