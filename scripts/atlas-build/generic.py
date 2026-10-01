"""Spec-driven loader: one reader for many datasets, so a new source needs a short JSON spec, not new code.

A spec (data/historical/specs/<id>.json) says where the data is, how to read it, and which fields mean what:

  {"id": "trajansgate", "src": "trajansgate", "public": true, "title": "...", "licence": "CC BY 4.0", "url": "...",
   "read": {"path": "trajans-gate/original/tgp.zip", "member": "tgp/tgp.gpkg", "layer": "tgp_sites", "crs": "EPSG:3035"},
   "fields": {"id": "fid", "name": "name", "kind": {"field": "type", "map": {"inn": "building"}, "default": "site"},
              "type": ["type"], "context": ["source"], "approx": {"field": "geocoding_accuracy", "values": ["200-500 m", ">500 m"]}},
   "dating": {"mode": "fields", "from": "start_date", "to": "end_date"}
         or {"mode": "text", "field": "dating"}          (one field holding "13th c.", "1250–1300", "c. 1400", "XIV. sz." …)
         or {"mode": "snapshot", "year": 1751}            (every record listed in one dated source)
         or {"mode": "envelope", "from": 1630, "to": 1655} (every record on one dated survey)
   "keep": {"field": "type", "values": [...]}            (optional filter)
  }

Readers: csv/tsv/txt, xlsx, geojson/json, gpkg, shp (in a zip or not), kml, ArcGIS FeatureServer layer ("read": {"arcgis": url}).
Positions are reprojected from the spec's CRS (or the .prj) with pyproj; lines and polygons give their representative point.

Dating follows the atlas rules: nothing is invented. A century is a window of that century (marked), "c." widens nothing but is
kept in the label, "?" marks the record uncertain in its label, an open end stays open, a missing date leaves the record undated.
"""
from __future__ import annotations

import csv
import glob
import io
import json
import math
import os
import re
import sqlite3
import tempfile
import urllib.request
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
RAW = os.path.join(HERE, '..', '..', 'data', 'historical', 'raw')
SPECS = os.path.join(HERE, '..', '..', 'data', 'historical', 'specs')

ROMAN = {'i': 1, 'ii': 2, 'iii': 3, 'iv': 4, 'v': 5, 'vi': 6, 'vii': 7, 'viii': 8, 'ix': 9, 'x': 10, 'xi': 11, 'xii': 12, 'xiii': 13,
         'xiv': 14, 'xv': 15, 'xvi': 16, 'xvii': 17, 'xviii': 18, 'xix': 19, 'xx': 20, 'xxi': 21}
PART = {'early': (0, 33), 'first half': (0, 50), '1st half': (0, 50), 'mid': (33, 66), 'middle': (33, 66), 'second half': (50, 100),
        '2nd half': (50, 100), 'late': (66, 100), 'end': (75, 100), 'beginning': (0, 25), 'start': (0, 25)}
BCE = re.compile(r'\b(bc|bce|b\.c\.|v\.\s?chr|av\.?\s?j\.?-?c|pr\.\s?n\.\s?l|př\.\s?n\.\s?l|i\.\s?e\.|до н\.\s?э)', re.I)


def _century_window(n, bce=False, part=None):
    a, b = (n - 1) * 100 + 1, n * 100
    if part:
        p0, p1 = PART[part]
        a, b = a + p0, a + p1 - 1 if p1 < 100 else b
    return (-b, -a) if bce else (a, b)


def _signed_range(text):
    """ArkeoGIS-style signed years: "-27", "-27:13", "451:475" (negative = BCE; there is no year 0)."""
    m = re.fullmatch(r'\s*(-?\d{1,4})\s*(?::\s*(-?\d{1,4}))?\s*', str(text or ''))
    if not m:
        return parse_dating(text)
    a = int(m.group(1))
    b = int(m.group(2)) if m.group(2) else a
    if a == 0 or b == 0 or a > b:
        return None
    return a, b, 'years' if a != b else 'year'


def kept(p, keeps) -> bool:
    """Whether a record passes a spec's "keep" filter(s): one dict or a list, each with "values" (exact allow-list),
    "notValues" (exact deny-list) or "notContaining" (deny any value containing one of these substrings)."""
    for k in ([keeps] if isinstance(keeps, dict) else keeps or []):
        v = str(p.get(k['field']))
        if 'values' in k and v not in k['values']:
            return False
        if v in k.get('notValues', []) or any(x in v for x in k.get('notContaining', [])):
            return False
    return True


_ROMAN_N = {'I': 1, 'II': 2, 'III': 3, 'IV': 4, 'V': 5, 'VI': 6, 'VII': 7, 'VIII': 8, 'IX': 9, 'X': 10, 'XI': 11, 'XII': 12, 'XIII': 13,
            'XIV': 14, 'XV': 15, 'XVI': 16, 'XVII': 17, 'XVIII': 18, 'XIX': 19, 'XX': 20, 'XXI': 21}


def lt_dating(text):
    """A Lithuanian heritage-register dating ("XIX a. pab.", "I t-metis – II t-mečio pr.", "XX a. 4 d-metis", "1895–1899 m.")
    in the English forms parse_dating reads; the parts keep their own meaning (pr. = beginning, vid. = middle, pab. = end,
    I/II p. = first/second half; t-metis = millennium). Returns the text unchanged where nothing matches."""
    if text is None:
        return None
    t = str(text)
    bc = bool(re.search(r'pr\.\s*Kr', t))
    t = re.sub(r'\b(?:po|pr\.)\s*Kr\.?', ' ', t)
    out = []
    # decades: "XX a. 4 d-metis" → 1930–1939
    for m in re.finditer(r'\b([IVX]+)\s*a\.\s*(\d)\s*d-?me(?:tis|čio)', t):
        c = _ROMAN_N.get(m.group(1))
        if c:
            y = (c - 1) * 100 + (int(m.group(2)) - 1) * 10  # the 4th decade of the 20th century = 1930s
            out.append(f'{y}-{y + 9}')
    t = re.sub(r'\b([IVX]+)\s*a\.\s*(\d)\s*d-?me(?:tis|čio)', ' ', t)
    part = {'pr.': 'early', 'pradž': 'early', 'pab.': 'late', 'pabaig': 'late', 'vid.': 'mid', 'I p.': 'first half', 'II p.': 'second half'}
    # millennia: "I t-metis", "II t-mečio pr." (beginning of the 2nd millennium) → the first / last two centuries of it
    for m in re.finditer(r'\b(I{1,3})\s*t-me(?:tis|čio|tyje)\s*(pr\.|pradž\w*|vid\.|pab\.|pabaig\w*|I p\.|II p\.)?', t):
        n = len(m.group(1)); a0, a1 = (n - 1) * 1000 + 1, n * 1000
        w = (m.group(2) or '').strip()
        if w.startswith('pr'):
            a1 = a0 + 199
        elif w.startswith('pab'):
            a0 = a1 - 199
        elif w.startswith('vid'):
            a0, a1 = a0 + 400, a0 + 599
        elif w == 'I p.':
            a1 = a0 + 499
        elif w == 'II p.':
            a0 = a0 + 500
        out.append(f'{a0 // 100 + 1}th century - {a1 // 100}th century' + (' BC' if bc else ''))  # as centuries: years below 100 are not read as years
    t = re.sub(r'\b(I{1,3})\s*t-me(?:tis|čio|tyje)\s*(pr\.|pradž\w*|vid\.|pab\.|pabaig\w*|I p\.|II p\.)?', ' ', t)
    # centuries: "XIX a. pab.", "XVI-XVIII a.", "XIX – XX a. I p."
    def cent(m):
        c = _ROMAN_N.get(m.group(1))
        if not c:
            return m.group(0)
        w = (m.group(2) or '').strip()
        p = next((v for k, v in part.items() if w.startswith(k)), '')
        return f' {p} {c}th century '.replace('  ', ' ')
    t = re.sub(r'\b([IVX]+)\s*a\.\s*(pr\.|pradž\w*|vid\.|pab\.|pabaig\w*|I p\.|II p\.)?', cent, t)
    t = re.sub(r'\b([IVX]+)\s*[-–]\s*(?=(?:early |late |mid |first half |second half )?\d+th century)', lambda m: f'{_ROMAN_N.get(m.group(1), 0)}th century - ', t)
    t = re.sub(r'\bm\.', ' ', t)
    return ' ; '.join(out + [t.strip()]) if out else t.strip()



def lv_dating(text):
    """A Latvian monument-list dating ("14.-15.gs.", "19. gs. I p.", "17.gs.b.", "1871.", "1927.-1961.") in the English forms
    parse_dating reads (gs. = century; s./sāk. = beginning, v./vid. = middle, b./beig. = end, I/II p. = first/second half)."""
    if text is None:
        return None
    t = str(text)
    low = t.lower()
    if re.search(r'bronz|akmens|neolīt|mezolīt|paleolīt', low):
        return None  # prehistoric periods: left undated rather than narrowed to the part that is mapped
    # the conventional Latvian archaeological periods (as centuries: years below 100 are not read as years)
    periods = [(r'agr\w* dzelzs laikmet\w*', '1st century - 4th century'), (r'vidēj\w* dzelzs laikmet\w*', '5th century - 8th century'),
               (r'vēl\w* dzelzs laikmet\w*', '9th century - 12th century'), (r'dzelzs laikmet\w*', '1st century - 12th century'),
               (r'viduslaik\w*', '1201-1561'), (r'jaun\w* laik\w*', '1561-1795')]
    for pat, rep in periods:
        low = re.sub(pat, f' {rep} ', low)
    t = low
    # quarters: "19. gs. 2.c." → 1826–1850
    t = re.sub(r'\b(\d{1,2})\.?\s*gs\.?\s*([1-4])\.\s*c\.', lambda m: f' {(int(m.group(1)) - 1) * 100 + (int(m.group(2)) - 1) * 25 + 1}-{(int(m.group(1)) - 1) * 100 + int(m.group(2)) * 25} ', t)
    t = re.sub(r'(\d{3,4})\.', r'\1', t)  # "1871." → 1871
    part = [(r'(?:s\.|sāk\w*\.?)', 'early'), (r'(?:b\.|beig\w*\.?)', 'late'), (r'(?:v\.|vid\w*\.?)', 'mid'),
            (r'(?:i\s*p\.|1\.\s*p\.|1\.\s*puse)', 'first half'), (r'(?:ii\s*p\.|2\.\s*p\.|2\.\s*puse)', 'second half')]

    def cent(m):
        w = (m.group(2) or '').strip()
        p = next((v for k, v in part if w and re.fullmatch(k, w)), '')
        return f' {p} {int(m.group(1))}th century '
    t = re.sub(r'\b(\d{1,2})\.?\s*gs\.?\s*(ii\s*p\.|i\s*p\.|[12]\.\s*p\.|[12]\.\s*puse|s\.|sāk\w*\.?|b\.|beig\w*\.?|v\.|vid\w*\.?)?', cent, t)
    t = re.sub(r'\b(\d{1,2})\.?\s*[-–/]\s*(?=\s*(?:early |late |mid |first half |second half )?\d+th century)', lambda m: f'{int(m.group(1))}th century - ', t)
    return re.sub(r'\s+', ' ', t).strip()



def parse_dating(text) -> tuple[int | None, int | None, str] | None:
    """(from, to, how) from a free-text dating, or None if it holds no date. Never more precise than the text."""
    if text is None:
        return None
    if isinstance(text, (int, float)) and not (isinstance(text, float) and math.isnan(text)):
        y = int(text)
        return (y, y, 'year') if -3000 <= y <= 2100 and y != 0 else None
    t = str(text).strip()
    if not t or t.lower() in ('nan', 'none', 'null', 'unknown', 'undetermined', 'neznámé', 'unbekannt', 'inconnu', '-', '?'):
        return None
    low = t.lower()
    bce = bool(BCE.search(low))
    # centuries: "13th c.", "13th century", "13. Jh.", "XIII. sz.", "XIIIe siècle", "XIII w.", "13. stol." — with an optional part
    cents = []
    CW = r'(?:c\b|c\.|cent|century|centuries|jh|jahrh|siècle|siecle|s\.|sec|secolo|siglo|sz|század|w\.|wiek|stol|století|stor|st\.|vek|век)'
    # century ranges: "5th-7th c.", "14-15 c.", "XII-XIII w."
    for m in re.finditer(r'\b(\d{1,2})(?:st|nd|rd|th|\.)?\s*[-–/]\s*(\d{1,2})(?:st|nd|rd|th|\.|e|er|ème)?\s*' + CW, low):
        a_, b_ = int(m.group(1)), int(m.group(2))
        if 1 <= a_ <= b_ <= 21:
            cents.append((_century_window(a_, bce)[0] if not bce else -b_ * 100, _century_window(b_, bce)[1] if not bce else -(a_ - 1) * 100 - 1))
    for m in re.finditer(r'\b([ivxl]{1,6})\.?\s*[-–/]\s*([ivxl]{1,6})\.?\s*' + CW, low):
        a_, b_ = ROMAN.get(m.group(1)), ROMAN.get(m.group(2))
        if a_ and b_ and a_ <= b_:
            cents.append((_century_window(a_, bce)[0], _century_window(b_, bce)[1]))
    for m in re.finditer(r'(?:(early|late|mid|middle|first half|1st half|second half|2nd half|end|beginning)(?:\s+of)?\s+(?:the\s+)?)?'
                         r'\b(\d{1,2})(?:st|nd|rd|th|\.|e|er|ème)?\s*(?:c\b|c\.|cent|century|centuries|jh|jahrh|siècle|s\.|sec|secolo|siglo|sz|század|w\.|wiek|stol|století|stor|st\.|vek|век)', low):
        part = (m.group(1) or '').strip() or None
        cents.append(_century_window(int(m.group(2)), bce, part if part in PART else None))
    for m in re.finditer(r'\b([ivxl]{1,6})\.?\s*(?:c\b|c\.|cent|century|jh|siècle|s\.|sec|secolo|siglo|sz|század|w\.|wiek|stol|st\.|vek|век|e\b|ème)', low):
        n = ROMAN.get(m.group(1))
        if n:
            cents.append(_century_window(n, bce))
    years = [int(y) for y in re.findall(r'(?<![\d.,])(\d{3,4})(?![\d.,])', t) if 100 <= int(y) <= 2100]
    # an abbreviated end year ("1852-62", "1863-4") belongs to the same century as its start
    for m in re.finditer(r'(?<![\d.,/-])(\d{4})\s*[-–/]\s*(\d{1,2})(?![\d.,])(?!\s*[-–/.]\s*\d)', t):  # not a full date (1801-12-05)
        a_, b_ = int(m.group(1)), m.group(2)
        end = int(str(a_)[:4 - len(b_)] + b_)
        if a_ < end <= 2100:
            years.append(end)
    # a decade ("1850s", "the 1850's"): its ten years
    for m in re.finditer(r'(?<![\d.,])(\d{3})0\'?s\b', t):
        years += [int(m.group(1) + '0'), int(m.group(1) + '9')]
    if bce:
        years = [-y for y in years]
    if cents and not years:
        return min(a for a, _ in cents), max(b for _, b in cents), 'century'
    if years:
        lo, hi = min(years), max(years)
        if cents:
            lo, hi = min(lo, min(a for a, _ in cents)), max(hi, max(b for _, b in cents))
        if re.search(r'\b(after|post|nach|après|po|od|from|since|seit)\b', low) and len(years) == 1:
            return lo, None, 'from year'
        if re.search(r'\b(before|ante|vor|avant|przed|do|until|bis)\b', low) and len(years) == 1 and not re.search(r'\bod\b', low):
            return None, hi, 'until year'
        return lo, hi, 'years' if lo != hi else 'year'
    return None


# ── readers ───────────────────────────────────────────────────────────────

def _raw(path):
    p = os.path.join(RAW, path)
    return p if os.path.exists(p) else None


def _raw_or_derive(read):
    p = _raw(read['path'])
    if p is None and read.get('derive'):
        # a file derived from originals in the vault (e.g. per-port attestation years): made by its script when missing
        import subprocess
        import sys as _sys
        subprocess.run([_sys.executable, os.path.join(os.path.dirname(__file__), '..', '..', read['derive'])], check=True)
        p = _raw(read['path'])
    return p


def _open_bytes(read):
    p = _raw_or_derive(read)
    if p is None:
        raise FileNotFoundError(read['path'])
    if read.get('member'):
        z = zipfile.ZipFile(p)
        return z.read(read['member']), read['member']
    return open(p, 'rb').read(), p


def _rows_table(b, name, read):
    e = name.lower().rsplit('.', 1)[-1]
    if e == 'xlsx':
        import openpyxl
        ws = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)[read['sheet']] if read.get('sheet') else \
            openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True).worksheets[0]
        rr = [list(r) for r in ws.iter_rows(values_only=True)]
        cols = [str(c or '') for c in rr[0]]
        return [dict(zip(cols, r)) for r in rr[1:]]
    t = b.decode(read.get('encoding', 'utf-8-sig'), 'replace')
    csv.field_size_limit(1 << 30)  # geometry columns (GeoJSON text) can be large
    if read.get('headerAfter'):  # a metadata block before the table (PANGAEA "/* … */"): the table starts on the next line
        t = t.split(read['headerAfter'], 1)[1].lstrip('\r\n')
    delim = read.get('delimiter') or csv.Sniffer().sniff(t[:20000], delimiters=',;\t|').delimiter
    return list(csv.DictReader(io.StringIO(t), delimiter=delim))


KEEP_GEOM = False


def _is_wgs84_degrees(crs):
    """True for plain WGS84 longitude/latitude. A projected CRS on the WGS84 datum (UTM, conic…) is not, even though its
    WKT names "WGS_1984" (a .prj may also carry a WGS84 vertical datum)."""
    s = str(crs)
    if re.fullmatch(r'\s*(EPSG:)?4326\s*', s, re.I):
        return True
    try:
        from pyproj import CRS
        c = CRS.from_user_input(s)
        return c.is_geographic and abs((c.datum.ellipsoid.semi_major_metre if c.datum and c.datum.ellipsoid else 6378137.0) - 6378137.0) < 1
    except Exception:  # noqa: BLE001
        return bool(re.search(r'GEOGCS\["?GCS_WGS_1984', s)) and 'PROJCS' not in s


def _geom_point(g):
    """(lon, lat) of a GeoJSON-like or shapely geometry: points as given, other shapes their representative point."""
    from shapely.geometry import shape
    s = g if hasattr(g, 'representative_point') else shape(g)
    if s.is_empty:
        return None
    p = s if s.geom_type == 'Point' else s.representative_point()
    return p.x, p.y


def _gpkg_rows(b, read):
    from shapely import wkb
    with tempfile.NamedTemporaryFile(suffix='.gpkg') as tf:
        tf.write(b)
        tf.flush()
        con = sqlite3.connect(tf.name)
        layer = read.get('layer') or con.execute("select table_name from gpkg_contents where data_type='features'").fetchone()[0]
        gcol, srs = con.execute('select column_name, srs_id from gpkg_geometry_columns where table_name=?', (layer,)).fetchone()
        cols = [r[1] for r in con.execute(f'pragma table_info("{layer}")')]
        out = []
        for r in con.execute(f'select * from "{layer}"'):
            d = dict(zip(cols, r))
            blob = d.pop(gcol, None)
            pt = None
            if blob:
                flags = blob[3]
                env = {0: 0, 1: 32, 2: 48, 3: 48, 4: 64}.get((flags >> 1) & 7, 0)
                g = wkb.loads(bytes(blob[8 + env:]))
                pt = _geom_point(g)
                if KEEP_GEOM:
                    from shapely.geometry import mapping
                    d['__geom'] = mapping(g)
            d['__pt'] = pt
            out.append(d)
        con.close()
    return out, read.get('crs') or (f'EPSG:{srs}' if srs and srs > 0 else None)


def _shp_rows(read):
    import shapefile
    p = _raw_or_derive(read)
    if p is None:
        raise FileNotFoundError(read['path'])
    if p.lower().endswith('.zip'):
        z = zipfile.ZipFile(p)
        base = read['member'][:-4]
        names = {n.lower(): n for n in z.namelist()}
        get = lambda ext: io.BytesIO(z.read(names[(base + ext).lower()]))  # noqa: E731
        prj = z.read(names[(base + '.prj').lower()]).decode('latin-1') if (base + '.prj').lower() in names else None
        r = shapefile.Reader(shp=get('.shp'), dbf=get('.dbf'), shx=get('.shx'), encoding=read.get('encoding', 'utf-8'))
    else:
        r = shapefile.Reader(p, encoding=read.get('encoding', 'utf-8'))
        prj = open(p[:-4] + '.prj').read() if os.path.exists(p[:-4] + '.prj') else None
    cols = [f[0] for f in r.fields[1:]]
    out = []
    for i, rec in enumerate(r.iterRecords()):
        d = dict(zip(cols, rec))
        try:
            shp = r.shape(i)
            d['__pt'] = _geom_point(shp.__geo_interface__) if shp.shapeType else None
            if KEEP_GEOM and shp.shapeType:
                d['__geom'] = shp.__geo_interface__
        except Exception:  # noqa: BLE001 — an empty or broken shape: the record is kept without a position
            d['__pt'] = None
        out.append(d)
    return out, read.get('crs') or prj


def _arcgis_rows(read):
    url = read['arcgis'].rstrip('/')
    out, offset = [], 0
    while True:
        u = f'{url}/query?where=1%3D1&outFields=*&outSR=4326&f=geojson&resultOffset={offset}&resultRecordCount=1000'
        with urllib.request.urlopen(urllib.request.Request(u, headers={'User-Agent': 'ShelfAtlasBuild/1.0'}), timeout=120) as r:
            d = json.load(r)
        fs = d.get('features') or []
        for f in fs:
            p = dict(f.get('properties') or {})
            p['__pt'] = _geom_point(f['geometry']) if f.get('geometry') else None
            out.append(p)
        if len(fs) < 1000 or not d.get('exceededTransferLimit', len(fs) == 1000):
            break
        offset += len(fs)
    return out, 'EPSG:4326'


def read_rows(spec):
    """[(props, (lon, lat) | None)] in WGS84."""
    read = spec['read']
    crs = read.get('crs')
    if read.get('arcgis'):
        rows, crs = _arcgis_rows(read)
    elif (read.get('member') or read['path']).lower().endswith('.shp'):
        rows, crs = _shp_rows(read)
    else:
        b, name = _open_bytes(read)
        e = name.lower().rsplit('.', 1)[-1]
        if e == 'gpkg':
            rows, crs = _gpkg_rows(b, read)
        elif e in ('geojson', 'json'):
            d = json.loads(b.decode('utf-8-sig'))
            rows = []
            crs = crs or ((d.get('crs') or {}).get('properties') or {}).get('name')
            for f in d.get('features') or []:
                p = dict(f.get('properties') or {})
                p['__pt'] = _geom_point(f['geometry']) if f.get('geometry') else None
                if KEEP_GEOM:
                    p['__geom'] = f.get('geometry')
                rows.append(p)
        elif e == 'kml':
            rows = []
            kml = re.sub(r'<!\[CDATA\[(.*?)\]\]>', lambda m: m.group(1).replace('<', '&lt;'), b.decode('utf-8', 'replace'), flags=re.S)
            for pm in re.findall(r'<Placemark\b.*?</Placemark>', kml, re.S):
                p = {k: v for k, v in re.findall(r'<SimpleData name="([^"]+)">([^<]*)<', pm)}
                p.update({k: v for k, v in re.findall(r'<Data name="([^"]+)">\s*<value>([^<]*)<', pm)})
                nm = re.search(r'<name>(.*?)</name>', pm, re.S)
                p.setdefault('name', nm.group(1).strip() if nm else None)
                for tag in ('begin', 'end', 'when'):
                    m = re.search(rf'<{tag}>([^<]+)</{tag}>', pm)
                    if m:
                        p[f'kml_{tag}'] = m.group(1)
                c = re.search(r'<coordinates>\s*([-\d.]+),([-\d.]+)', pm)
                p['__pt'] = (float(c.group(1)), float(c.group(2))) if c else None
                rows.append(p)
            crs = 'EPSG:4326'
        else:
            rows = _rows_table(b, name, read)
            if read.get('join'):
                # attach a second table's fields (e.g. a gazetteer's coordinates) by key
                j = read['join']
                jb, jn = _open_bytes(j)
                jt = {str(r.get(j['key'])).strip(): r for r in _rows_table(jb, jn, j)}
                for p in rows:
                    other = jt.get(str(p.get(j['on'])).strip())
                    if other:
                        for k, v in other.items():
                            p.setdefault(k, v)
            lon_f, lat_f = read.get('lon'), read.get('lat')
            for p in rows:
                try:
                    p['__pt'] = (float(str(p[lon_f]).replace(',', '.')), float(str(p[lat_f]).replace(',', '.')))
                except (KeyError, TypeError, ValueError):
                    p['__pt'] = None
                if p['__pt'] is None and read.get('geojsonCol') and p.get(read['geojsonCol']):
                    # an area (a concession, a parish) given as GeoJSON text: its representative point, marked approximate
                    try:
                        g = json.loads(p[read['geojsonCol']])
                        p['__pt'] = _geom_point(g) if isinstance(g, dict) and g.get('type') else None
                        p['__approx'] = True
                        if KEEP_GEOM:
                            p['__geom'] = g
                    except (ValueError, TypeError, KeyError, AttributeError):
                        pass
                    p.pop(read['geojsonCol'], None)  # large text, not needed after this
                if p['__pt'] is None and read.get('wkbHex') and p.get(read['wkbHex']):
                    # an area (island, region) given as hex WKB: its representative point, marked approximate
                    try:
                        from shapely import wkb
                        p['__pt'] = _geom_point(wkb.loads(bytes.fromhex(str(p[read['wkbHex']]).strip())))
                        p['__approx'] = True
                    except Exception:  # noqa: BLE001
                        pass
    if read.get('pointFields') and rows:
        # the source's own WGS84 coordinate fields win over its (projected) geometry
        lon_k, lat_k = read['pointFields']
        for p in rows:
            try:
                p['__pt'] = (float(str(p[lon_k]).replace(',', '.')), float(str(p[lat_k]).replace(',', '.')))
                p['__wgs'] = True
            except (KeyError, TypeError, ValueError):
                pass
    if read.get('coalesce') and rows:
        # the first field with a value (e.g. an exact year before a period class) into one field
        co = read['coalesce']
        for p in rows:
            p[co['into']] = next((p.get(k) for k in co['fields'] if _val(p.get(k)) is not None), None)
    if read.get('translateDating') and rows:
        # a dating field in another language's conventions, translated into __dating (the original stays as it is)
        tr = {'lt': lt_dating, 'lv': lv_dating}[read['translateDating']['lang']]
        for p in rows:
            p['__dating'] = tr(p.get(read['translateDating']['field']))
    if read.get('fixMojibake') and rows:
        # UTF-8 text that was decoded as Latin-1 once ("maÃ§onnées" → "maçonnées"); left as is when it does not round-trip
        for p in rows:
            for k, v in p.items():
                if isinstance(v, str) and ('Ã' in v or 'Â' in v):
                    try:
                        p[k] = v.encode('latin-1').decode('utf-8')
                    except (UnicodeEncodeError, UnicodeDecodeError):
                        pass
    if read.get('latLonField') and rows:
        # one "lat, lon" text field (e.g. a commune's centroid); "approx": the position is the place's, not the feature's
        for p in rows:
            m = re.fullmatch(r'\s*(-?\d+(?:\.\d+)?)\s*,\s*(-?\d+(?:\.\d+)?)\s*', str(p.get(read['latLonField']) or ''))
            if m:
                p['__pt'] = (float(m.group(2)), float(m.group(1)))
                p['__wgs'] = True
                if read.get('approx'):
                    p['__approx'] = True
    if crs and not _is_wgs84_degrees(crs):
        from pyproj import CRS, Transformer
        tr = Transformer.from_crs(CRS.from_user_input(crs), CRS.from_epsg(4326), always_xy=True)
        for p in rows:
            if p.get('__pt') and not p.get('__wgs'):
                p['__pt'] = tr.transform(*p['__pt'])
            if p.get('__geom'):
                from shapely.geometry import mapping, shape
                from shapely.ops import transform
                p['__geom'] = mapping(transform(lambda x, y, z=None: tr.transform(x, y), shape(p['__geom'])))
    return rows


# ── spec → register records ───────────────────────────────────────────────

NULLISH = {'', 'no comment', 'n/a', 'na', 'none', 'null', 'unclear', 'unknown', '-', 'nan'}


def _val(v):
    return None if v is None or str(v).strip().lower() in NULLISH else v


def _field(p, f):
    if isinstance(f, list):
        return ' · '.join(str(p.get(x)) for x in f if _val(p.get(x)) is not None)
    return _val(p.get(f)) if f else None


def parts(spec):
    """A spec with "parts" (several files, each with its own read/dating) as one spec per part."""
    if not spec.get('parts'):
        return [spec]
    return [{**spec, **pt, 'fields': {**spec.get('fields', {}), **pt.get('fields', {})}} for pt in spec['parts']]


def records(spec):
    """Register records (registers.py schema) from a spec; and counts of what was left out and why."""
    if spec.get('parts'):
        from collections import Counter
        out, skipped, seen = [], Counter(), set()
        for k, sp in enumerate(parts({k: v for k, v in spec.items()})):
            sp.pop('parts', None)
            rs, sk = records(sp)
            skipped.update(sk)
            for r in rs:
                key = (r['name'], r['lon'], r['lat'], r.get('env'), r.get('snap'))
                if spec.get('dedupe') and key in seen:
                    skipped['same record in another part'] += 1
                    continue
                seen.add(key)
                r['id'] = f"{k}:{r['id']}"
                out.append(r)
        return out, dict(skipped)
    from collections import Counter
    f, dt = spec['fields'], spec['dating']
    out, skipped = [], Counter()
    excl = set()
    if spec.get('excludeIdsFrom'):
        # records another dataset already shows (e.g. a public subset of the same database): left out here, not drawn twice
        ex = spec['excludeIdsFrom']
        excl = {str(r.get(ex['field'])).strip() for r in read_rows({'read': ex['read']})}
    for i, p in enumerate(read_rows(spec)):
        if excl and str(p.get(spec['fields'].get('id'))).strip() in excl:
            skipped['already shown by ' + spec['excludeIdsFrom'].get('label', 'another dataset')] += 1
            continue
        if not kept(p, spec.get('keep')):
            skipped['outside the spec filter'] += 1
            continue
        pt = p.get('__pt')
        if not pt or not (-180 <= pt[0] <= 180 and -90 <= pt[1] <= 90) or (pt[0] == 0 and pt[1] == 0):
            skipped['no usable position'] += 1
            continue
        name = str(_field(p, f.get('name')) or '').strip()
        if name and f.get('nameDropParens'):
            name = re.sub(r'\s*\([^)]*\)?', '', name).strip() or name  # "ENCLOSURE (A.P. SITE: unlocated)" → "ENCLOSURE"
        generic_name = bool(f.get('nameSplitRequired')) and f.get('nameSplit', '') not in name
        if name and f.get('nameSplit'):
            # "Abae, Achaea" → "Abae" (the rest is the source's context); index -1: "FORTIFIED OUTCROP: DUNSHAMMER" → "DUNSHAMMER"
            name = name.split(f['nameSplit'])[f.get('nameSplitIndex', 0)].strip()
        if name and f.get('nameTitleCase'):
            # capitals-only words → capitalised ("GIANT'S SCONCE or DUNCEITHIRN" → "Giant's Sconce or Dunceithirn")
            name = ' '.join(w[0] + w[1:].lower() if w.isupper() and len(w) > 1 else w for w in name.split())
        if not name and f.get('nameFallback'):
            name = str(_field(p, f['nameFallback']) or '').strip()
        kf = f.get('kind') or {}
        kv = str(p.get(kf.get('field')) if p.get(kf.get('field')) is not None else '').strip().lower() if isinstance(kf, dict) else ''
        if isinstance(kf, dict) and kf.get('firstToken'):
            kv = kv.split(';')[0].strip()
        kind = (kf.get('map') or {}).get(kv, kf.get('default', 'site')) if isinstance(kf, dict) else kf
        tmap = f.get('typeMap') or {}
        tparts = [str(tmap.get(k, {}).get(str(p.get(k)), p.get(k))) for k in ([f['type']] if isinstance(f.get('type'), str) else f.get('type') or []) + (f.get('typeExtra') or [])
                  if _val(p.get(k)) is not None]
        rec = {'src': spec['src'], 'id': str(p.get(f.get('id')) if f.get('id') else i), 'name': (name or spec.get('unnamed', 'Unnamed site'))[:80],
               'kind': kind, 'lon': round(pt[0], 5), 'lat': round(pt[1], 5), 'ty': (' · '.join(tparts) or kind)[:120],
               'ctx': [str(_field(p, c)) for c in f.get('context', []) if _field(p, c)][:2], 'names': []}
        if f.get('altNames'):
            # variant spellings the source records for the same place ("a coruña|corvigna|la corugna")
            an = f['altNames']
            alts = re.split(an.get('sep', r'\|'), str(p.get(an['field']) or ''))
            rec['names'] = [(a, None, None, '') for a in sorted({a.strip() for a in alts if a.strip() and a.strip().lower() != name.lower()})[:12]]
        # survey codes as names ("SGNAS SITE 018", genericPattern) and type-only names are drawn but not indexed as places
        if f.get('generic') or generic_name or (f.get('genericIfEquals') and name.strip().lower() == str(p.get(f['genericIfEquals']) or '').strip().lower()) or (f.get('genericPattern') and re.fullmatch(f['genericPattern'], name)) or (f.get('nameSplitRequired') and name[:1].islower()):  # "reputedly site of a massacre…" is a note, not a name
            rec['generic'] = 1  # the name is only the monument type ("Rath"): drawn and searchable by type, not a place name
        ap = f.get('approx')
        rec['precise'] = not (ap and str(p.get(ap['field'])) in ap['values']) and not p.get('__approx')
        uncertain = False
        dt = spec['dating']  # per record (a fallback below may switch this record's mode)
        if dt['mode'] == 'textOrPeriods':
            own = str(p.get(dt['field']) or '')
            pr = parse_dating(own) if not re.fullmatch(r'[A-Za-z\-]+', own.strip()) else None
            if pr and pr[0] is not None:
                rec['env'] = (pr[0], pr[1])
                rec['per'] = f"{own} ({dt.get('ownLabel', 'record dating')})"[:120]
                out.append(rec)
                continue
            dt = dict(dt, mode='periods', field=dt.get('periodField', dt['field']))
        if dt['mode'] == 'flags':
            # one yes/no column per period (per_ROM = 1, per_VIK = 1 …): the flagged periods, then dated like named periods
            on = {str(v) for v in dt.get('true', ['1', 'True', 'true', 'J', 'x'])}
            p['__flags'] = ';'.join(name for col, name in dt['flags'].items() if str(p.get(col)) in on)
            dt = dict(dt, mode='periods', field='__flags', sep=';')
        if dt['mode'] == 'periods':
            # Named periods (one field, separated), each mapped by the spec's own period table (a stated convention, e.g.
            # standard Egyptian chronology); separate periods stay separate phases (merged only when overlapping or within 50 years).
            raw_label = str(p.get(dt['field']) or '')
            spans = []
            for part in re.split(dt.get('sep', ';'), raw_label):
                low = re.sub(r'\s+', ' ', part.lower())
                unc = '?' in low
                # every period name of the table found in the text, longest first, without overlaps
                found, taken = [], []
                for key in sorted(dt['table'], key=len, reverse=True):
                    for m in re.finditer(rf'(?<!\w){re.escape(key)}(?!\w)', low):  # keys may end in '.' ("e.christ.")
                        if not any(m.start() < e and s_ < m.end() for s_, e in taken):
                            taken.append((m.start(), m.end()))
                            found.append(key)
                rest = re.sub(r'[\W_]+', ' ', ''.join(ch if not any(s_ <= i < e for s_, e in taken) else ' ' for i, ch in enumerate(low))).strip()
                if rest and rest not in dt.get('ignore', []):
                    skipped[f'text not in the spec period table: {rest}'] += 1
                for key in found:
                    rng = dt['table'][key]
                    if rng[1] < dt.get('min', 1) or rng[0] > dt.get('max', 1914):
                        continue
                    spans.append((rng[0], rng[1], key.capitalize() + (' (uncertain)' if unc else '')))
            spans.sort()
            phases = []
            for a_, b_, n_ in spans:
                if phases and a_ <= phases[-1][1] + 50:
                    phases[-1] = (phases[-1][0], max(phases[-1][1], b_), phases[-1][2] + [n_])
                else:
                    phases.append((a_, b_, [n_]))
            if not phases:
                skipped['no period in the atlas range'] += 1
                if not spec.get('keepUndated'):
                    continue
                out.append(rec)
                continue
            for k, (a_, b_, ns) in enumerate(phases):
                r2 = dict(rec, id=f"{rec['id']}:{k}", env=(a_, b_), per=f"{'; '.join(ns)} ({dt.get('label', 'period table in the spec')}: {a_}–{b_})"[:120])
                out.append(r2)
            continue
        if dt['mode'] == 'firstAttested':
            # the earliest dated attestation in the field: the place is shown from then on (first attestation ≠ founding)
            raw_label = str(p.get(dt['field']) or '')
            cands = []
            for chunk in re.split(dt.get('sep', r'\n|;'), raw_label):
                pr = parse_dating(chunk)
                if pr and pr[0] is not None:
                    cands.append((pr[0], pr[1], chunk.strip()))
            if not cands:
                skipped['no dated attestation'] += 1
                if not spec.get('keepUndated'):
                    continue
                out.append(rec)
                continue
            lo, hi, chunk = min(cands)
            rec['env'] = (lo, None)
            rec['fa'] = 1
            rec['per'] = f"first attested {lo}{'–' + str(hi) if hi and hi != lo else ''} ({chunk[:70]})"[:120]
            out.append(rec)
            continue
        if dt['mode'] == 'attestations':
            # one snapshot per dated source the record is listed in (an itinerary, a pilgrimage guide…): the place is
            # shown at each attestation year only, nothing is inferred between or around them
            raw_label = str(p.get(dt['field']) or '')
            years = sorted({int(y) for y in re.findall(dt.get('yearRegex', r'(?<!\d)(\d{3,4})(?!\d)'), raw_label)
                            if dt.get('min', 1) <= int(y) <= dt.get('max', 1914)})
            if not years:
                skipped['no dated attestation'] += 1
                continue
            for y in years:
                out.append(dict(rec, id=f"{rec['id']}:{y}", snap=y,
                                per=dt.get('label', 'listed in a source of {year}').format(year=y, all=', '.join(map(str, years)))[:120]))
            continue
        if dt['mode'] == 'snapshot':
            rec['snap'] = dt['year']
            rec['per'] = dt.get('label') or f"listed in {dt['year']}"
        elif dt['mode'] == 'envelope':
            rec['env'] = (dt['from'], dt['to'])
            rec['per'] = dt.get('label') or f"{dt['from']}–{dt['to']} (source's survey window)"
        else:
            if dt['mode'] == 'fields':
                a, b = p.get(dt['from']), p.get(dt.get('to'))
                opens = {str(v) for v in dt.get('openValues', [])}
                if str(a) in opens:
                    a = None  # the source's code for "unknown"
                b_open = b not in (None, '') and str(b) in opens
                pa = _signed_range(a) if dt.get('signedYears') else parse_dating(a)
                pb = None if b in (None, '') or b_open else (_signed_range(b) if dt.get('signedYears') else parse_dating(b))
                if pa is None and dt.get('centuryField') and str(p.get(dt['centuryField']) or '').strip().isdigit():
                    # only the century is given ("11"): that century's window
                    n = int(str(p[dt['centuryField']]).strip())
                    if 1 <= n <= 20:
                        pa = (*_century_window(n, False), 'century')
                        a = f'{n}th century'
                if dt.get('midCenturyAsCentury') and pa and pa[0] == pa[1] and pa[0] % 100 == 50:
                    # "1350" in a source that encodes "14th century" as its midpoint: the century, not that year
                    pa = (*_century_window(pa[0] // 100 + 1, False), 'century')
                    a = f'{a} (a century midpoint: {pa[0]}–{pa[1]})'
                raw_label = ' – '.join([str(a)] if a not in (None, '') else []) + (' – (end unknown or continuing)' if b_open else f' – {b}' if b not in (None, '') else '')
                uncertain = '?' in raw_label
                lo = pa[0] if pa else None
                # a stated end that is "unknown/continuing" leaves the end open; no end field at all keeps the start's own span
                hi = None if b_open else pb[1] if pb else (pa[1] if pa and b in (None, '') else None)
                if pa and pa[2] == 'from year':
                    hi = pb[1] if pb else None
            else:
                raw_label = str(p.get(dt['field']) or '')
                uncertain = '?' in raw_label
                pr = parse_dating(raw_label)
                lo, hi = (pr[0], pr[1]) if pr else (None, None)
            if lo is None and hi is None and dt.get('fallback'):
                fb = dt['fallback']
                lo, hi = fb['from'], fb['to']
                raw_label = fb['label']
                skipped['no own date: part period used (' + fb['label'][:40] + ')'] += 1
            if lo is None and hi is None:
                skipped['no date in the record'] += 1
                if not spec.get('keepUndated'):
                    continue
            elif lo is None:
                if spec.get('dropUndated'):
                    skipped['only an end date (left out)'] += 1
                    continue
                skipped['only an end date (kept undated)'] += 1
                lo = hi = None
            if lo is not None:
                if hi is not None and hi < lo:
                    skipped['end before start (dates kept as the start only)'] += 1
                    hi = None
                rec['env'] = (lo, hi)
            rec['per'] = f"{raw_label}{' (uncertain in the source)' if uncertain and '?' not in raw_label else ''} ({spec.get('datingLabel', 'source dating')})"[:120]
        out.append(rec)
    if spec.get('groupBy'):
        out = _group_phases(out, spec, skipped)
    return out, dict(skipped)


def _group_phases(recs, spec, skipped):
    """Records of one place (spec "groupBy": the record name + position) merged into phases: dated evidence that overlaps or
    lies within 50 years is one phase; a gap stays a gap. The labels of the merged evidence are kept (counted)."""
    from collections import defaultdict
    groups = defaultdict(list)
    for r in recs:
        groups[(r['name'], r['lon'], r['lat'])].append(r)
    out = []
    for (name, lon, lat), rs in groups.items():
        dated = sorted((r for r in rs if r.get('env') and r['env'][0] is not None), key=lambda r: r['env'][0])
        if spec.get('groupFirstAttested') and dated:
            # the place from its earliest dated record on (first attestation ≠ founding); later records are counted, not used as ends
            first, last = dated[0]['env'][0], max(r['env'][1] or r['env'][0] for r in dated)
            base = dated[0]
            out.append(dict(base, id=f"{base['id']}:fa", env=(first, None), fa=1, ty=spec.get('groupType', base['ty']),
                            per=f"first attested {first} ({spec.get('groupNoun', 'record')}; {len(dated)} dated mention{'s' if len(dated) > 1 else ''} {first}–{last})"[:120]))
            continue
        phases = []
        for r in dated:
            a_, b_ = r['env'][0], r['env'][1] if r['env'][1] is not None else r['env'][0]
            # an end the source leaves open ("undetermined") stays open for the whole phase — never closed, never filled in
            open_end = r['env'][1] is None and bool(spec.get('dating', {}).get('openValues'))
            if phases and (phases[-1]['open'] or a_ <= phases[-1]['to'] + 50):
                phases[-1]['to'] = max(phases[-1]['to'], b_)
                phases[-1]['n'] += 1
                phases[-1]['labels'].add(r['ty'])
                phases[-1]['open'] = phases[-1]['open'] or open_end
            else:
                phases.append({'from': a_, 'to': b_, 'n': 1, 'labels': {r['ty']}, 'base': r, 'open': open_end})
        for k, ph in enumerate(phases):
            base = ph['base']
            end = None if ph['open'] else ph['to']
            out.append(dict(base, id=f"{base['id']}:{k}", env=(ph['from'], end),
                            per=f"{spec.get('groupVerb', 'attested by')} {ph['n']} {spec.get('groupNoun', 'dated record')}{'s' if ph['n'] > 1 else ''} ({', '.join(sorted(ph['labels']))[:60]}): {ph['from']}–{ph['to'] if end is not None else '(end undetermined)'}"[:120],
                            ty=spec.get('groupType', base['ty'])))
        if len(rs) > len(dated):
            skipped['undated evidence (not used)'] += len(rs) - len(dated)
    return out


def specs(public=None, geometry='points'):
    out = []
    for p in sorted(glob.glob(os.path.join(SPECS, '*.json'))):
        s = json.load(open(p, encoding='utf-8'))
        if geometry and s.get('geometry', 'points') != geometry:
            continue
        if s.get('enabled', True) and (public is None or bool(s.get('public')) == public):
            out.append(s)
    return out


APP_MODULE = os.path.join(HERE, '..', '..', 'src', 'atlas', 'spec-datasets.ts')


def write_app_module(loaded):
    """src/atlas/spec-datasets.ts: each spec's credit, period, extent and map layers, for the app's registries.
    `loaded` is {src: records} from this build; coverage and box come from the records themselves."""
    def q(s):
        return json.dumps(s, ensure_ascii=False)
    lines = ['// Generated by scripts/atlas-build/generic.py from data/historical/specs/*.json during the world build. Do not edit by hand.',
             'export const SPEC_DATASETS = {']
    for s in specs(geometry=None):
        recs = loaded.get(s['src']) or (line_features(s) if s.get('geometry') == 'lines' else area_features(s) if s.get('geometry') == 'polygons' else [])
        if s.get('geometry') in ('lines', 'polygons'):
            ys = [y for _, p, _ in recs for y in (p.get('ef'), p.get('et')) if y is not None]
            from shapely.geometry import shape
            bs = [shape(g).bounds for g, _, _ in recs]
            lons, lats = [b[0] for b in bs] + [b[2] for b in bs], [b[1] for b in bs] + [b[3] for b in bs]
        else:
            ys = [y for r in recs for y in ((r.get('env') or (None, None)) + ((r['snap'], r['snap']) if r.get('snap') else ())) if y is not None]
            lons, lats = [r['lon'] for r in recs], [r['lat'] for r in recs]
        cov = [min(ys), max(ys)] if ys else [1, 1914]
        box = [round(min(lons), 1), round(min(lats), 1), round(max(lons), 1), round(max(lats), 1)] if recs else [-25, 27, 62, 72]
        lines.append(f"  {s['src']}: {{ name: {q(s['title'])}, license: {q(s['licence'])}, url: {q(s['url'])}, coverage: [{cov[0]}, {cov[1]}] as [number, number], "
                     f"core: [{s.get('core', cov)[0]}, {s.get('core', cov)[1]}] as [number, number], box: [{', '.join(map(str, box))}] as [number, number, number, number], "
                     f"describe: {q(s.get('describe') or s['title'])}, public: {'true' if s.get('public') else 'false'}, layers: {q(s.get('layers') or ['medieval-archaeology'])} as string[] }},")
    lines += ['} as const;', 'export type SpecDatasetId = keyof typeof SPEC_DATASETS;', '']
    open(APP_MODULE, 'w', encoding='utf-8').write('\n'.join(lines))


def line_features(spec):
    """Dated line features (roads, routes, waterways) for a tile set: (GeoJSON geometry, props, min zoom).
    Props follow the dated-roads layer: ef/et (the part's period), n (name), k (type), u (1 = position uncertain), src, per."""
    global KEEP_GEOM
    out = []
    for k, sp in enumerate(parts(spec)):
        KEEP_GEOM = True
        try:
            rows = read_rows(sp)
        finally:
            KEEP_GEOM = False
        f, dt = sp.get('fields', {}), sp['dating']
        unc = f.get('uncertain') or {}
        for i, p in enumerate(rows):
            if not kept(p, sp.get('keep')):
                continue
            g = p.get('__geom')
            if not g or g.get('type') not in ('LineString', 'MultiLineString'):
                continue
            pr = {'i': f"{spec['src']}:{k}:{p.get(f.get('id')) if f.get('id') else i}", 'src': spec['src'], 'k': str(f.get('lineKind') or _field(p, f.get('type')) or 'road')[:40],
                  'u': 1 if unc and str(p.get(unc['field'])) in unc['values'] else 0, 'per': dt.get('label', '')[:80]}
            nm = _field(p, f.get('name'))
            if nm:
                pr['n'] = str(nm)[:60]
            for extra in f.get('keep', []):
                if p.get(extra) not in (None, ''):
                    pr[extra[:12]] = str(p[extra])[:60]
            if dt['mode'] == 'envelope':
                pr['ef'], pr['et'] = dt['from'], dt['to']
            elif dt['mode'] == 'snapshot':
                pr['ef'] = pr['et'] = dt['year']
                pr['sn'] = 1
            else:
                continue  # lines are drawn only with their part's own dating
            out.append((g, pr, f.get('minzoom', 5)))
    return out


def area_features(spec):
    """Dated areas (historical administrative or political units) for a tile set: (GeoJSON geometry, props, min zoom).
    With "dissolve": {"field": ..., "level": ...} the source polygons are merged per value of that field (e.g. localities →
    counties); several dissolve levels may be given. Props: n (unit name), lv (level label), ef/et or sn (the source's dating), src, per."""
    global KEEP_GEOM
    from shapely.geometry import mapping, shape
    from shapely.ops import unary_union
    out = []
    for k, sp in enumerate(parts(spec)):
        KEEP_GEOM = True
        try:
            rows = read_rows(sp)
        finally:
            KEEP_GEOM = False
        rows = [p for p in rows if kept(p, sp.get('keep'))]
        dt = sp['dating']
        date = {'ef': dt['year'], 'et': dt['year'], 'sn': 1} if dt['mode'] == 'snapshot' else \
            {'ef': dt['from'], 'et': dt['to']} if dt['mode'] == 'envelope' else None  # 'fields': each unit's own dates, below
        names = sp.get('names') or {}
        for lv in sp.get('dissolve') or []:
            # "constant": every polygon of the part is one unit of that name (e.g. a state's extent drawn in several pieces)
            fields = [] if lv.get('constant') else lv['field'] if isinstance(lv['field'], list) else [lv['field']]  # a list: units keyed by several fields (county + hundred)
            groups, labels, spans, notes = {}, {}, {}, {}
            for p in rows:
                g = p.get('__geom')
                if not g or any(p.get(x) in (None, '') for x in fields):
                    continue
                # contested/divided units: the first named holder (the source lists them ';'-separated)
                v = lv['constant'] if lv.get('constant') else '|'.join(str(p.get(x)).split(';')[0].strip() for x in fields)
                groups.setdefault(v, []).append(shape(g).buffer(0))
                labels.setdefault(v, v if lv.get('constant') else str(p.get(lv.get('nameField', fields[-1]))).split(';')[0].strip())
                if lv.get('note') and _val(p.get(lv['note']['field'])) is not None:  # a figure the source gives for the unit ("81,664 inhabitants")
                    nv = p[lv['note']['field']]
                    notes.setdefault(v, f"{int(float(nv)):,}" if re.fullmatch(r'\d+(\.0+)?', str(nv).strip()) else str(nv).strip())
                if date is None:
                    if dt.get('from'):  # separate start / end fields ("0725/01/01", "1794/12/31")
                        a_, b_ = parse_dating(p.get(dt['from'])), parse_dating(p.get(dt.get('to')))
                        pd = (a_[0], b_[1] if b_ else a_[1], 'fields') if a_ and a_[0] is not None else None
                    else:
                        pd = parse_dating(p.get(dt['field']))
                    if pd and pd[0] is not None:
                        a_, b_ = spans.get(v, (pd[0], pd[1] or pd[0]))
                        spans[v] = (min(a_, pd[0]), max(b_, pd[1] or pd[0]))
            for v, geoms in groups.items():
                if date is None and v not in spans:
                    continue  # no date of its own: not drawn
                u = unary_union(geoms).simplify(lv.get('simplify', 0.0005), preserve_topology=True)
                if u.is_empty:
                    continue
                d_ = date or {'ef': spans[v][0], 'et': spans[v][1], **({'sn': 1} if dt.get('snapshot') and spans[v][0] == spans[v][1] else {})}
                key = '|'.join(fields) or 'unit'
                lab = labels[v] if len(fields) > 1 or lv.get('nameField') else v
                per = dt.get('label', '') if date else f"{dt.get('label', 'surveyed')} {d_['ef']}–{d_['et']}" if d_['ef'] != d_['et'] else f"{dt.get('label', 'surveyed')} {d_['ef']}"
                if v in notes:
                    per = f"{per}; {notes[v]} {lv['note'].get('label', '')}".strip()
                out.append((mapping(u), {'i': f"{spec['src']}:{key}:{v}" + (f':{k}' if k else ''), 'n': names.get(fields[0] if len(fields) == 1 else key, {}).get(lab, lab)[:60],
                                         'lv': lv['level'], 'src': spec['src'], 'per': per[:80], **d_}, lv.get('minzoom', 4)))
    return out
